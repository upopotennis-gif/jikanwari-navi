#!/usr/bin/env python3
"""教員用「時間割ナビ（教員）」のデータを作り、合言葉で暗号化して ../jikanwari-navi-staff/data.enc.js に書き出す。

  先生の時間割 … source/jikanwari.xlsx「職員版配布用」＋ source/tokubetsu/*.pdf（特別時間割・教員版）
  教室        … 「教室配当」で同じ表記の授業が使う教室
  予定        … 予定.csv（生徒用と同じ）＋ source/gyouji_MM.pdf（月間行事予定の全部の欄）
  合言葉      … source/staff_password.txt（公開しない）

先に build.py を動かして data.js を最新にしておくこと。
"""
import base64, csv, json, os, re, sys, datetime, pathlib
from collections import Counter
import openpyxl
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
from cryptography.hazmat.primitives import hashes

from tokubetsu import parse_pdf
from build import LHR_ROOM
from kakutei import parse_rooms
from invert import parse_label
from build import BLOCKS
import gyouji_pdf

ROOT = pathlib.Path(__file__).parent
OUT = ROOT.parent / 'jikanwari-navi-staff'
FW = str.maketrans('０１２３４５６７８９ＡＢＣＤＥＦＬＨＲ', '0123456789ABCDEFLHR')
# 名前の読みかえ（特別時間割の表記にそろえる）。名前が入るので source/ に置く（公開しない）
_rn = pathlib.Path(__file__).parent / 'source' / 'staff_rename.json'
RENAME = json.loads(_rn.read_text(encoding='utf-8')) if _rn.exists() else {}
ITER = 200_000


def norm(s):
    return re.sub(r'[\s①②③]', '', str(s)).translate(FW)


def filled(ws):
    g = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                g[(c.row, c.column)] = c.value
    for m in ws.merged_cells.ranges:
        v = g.get((m.min_row, m.min_col))
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                g[(r, c)] = v
    return g


def main():
    student = json.loads((ROOT / 'data.js').read_text(encoding='utf-8')[len('window.DATA = '):-2])
    wb = openpyxl.load_workbook(ROOT / 'source' / 'jikanwari.xlsx', data_only=True)

    # ---- 先生の一覧とふだんの時間割
    ws = wb['職員版配布用']
    teachers, base = [], {}
    kyoka = ''
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, 1).value:
            kyoka = str(ws.cell(r, 1).value).strip()
        name = ws.cell(r, 5).value
        if not name:
            continue
        name = RENAME.get(str(name).strip(), str(name).strip())
        teachers.append({'name': name, 'kyoka': kyoka,
                         'bunsho': str(ws.cell(r, 2).value or '').strip(),
                         'gakunen': str(ws.cell(r, 3).value or '').strip()})
        base[name] = [str(ws.cell(r, 6 + i).value).strip() if ws.cell(r, 6 + i).value else '' for i in range(35)]

    # ---- 教室配当: その時間にその表記の授業がどの教室か
    wr = wb['教室配当']
    g = filled(wr)
    rooms_list, floor = [], ''
    by_slot = {}       # slot -> {norm表記: [教室]}
    used = {}          # slot -> set(教室)
    for r in range(6, wr.max_row + 1):
        if g.get((r, 1)):
            floor = re.sub(r'\s', '', str(g[(r, 1)]))
        room = re.sub(r'\s', '', str(g.get((r, 2)) or '')).replace('Ｈ', 'H').replace('Ｂ', 'B')
        if not room:
            continue
        rooms_list.append({'room': room, 'floor': floor})
        for i in range(35):
            lab = g.get((r, 3 + i))
            if lab:
                by_slot.setdefault(i, {}).setdefault(norm(lab), []).append(room)
                used.setdefault(i, set()).add(room)
    any_slot = {}
    for i, d in by_slot.items():
        for k, v in d.items():
            any_slot.setdefault(k, Counter())[tuple(sorted(set(v)))] += 1

    def room_of(label, slot=None):
        if not label or '\n' not in label and not re.search(r'[１２３][Ａ-Ｆ]', label):
            return [], False
        k = norm(label)
        if slot is not None and k in by_slot.get(slot, {}):
            return sorted(set(by_slot[slot][k])), False
        if k in any_slot:
            return list(any_slot[k].most_common(1)[0][0]), True
        return [], False

    def from_class(label, slots_of):
        """生徒用データ（クラスごとの教室）から引く。slots_of(クラスid) → そのコマのエントリ"""
        out = []
        subj = norm(label.split('\n')[0]) if label else ''
        for cid, kind, val in parse_label(label):
            e = slots_of(cid)
            if not e:
                continue
            if 'block' in e:
                for o in BLOCKS.get(e['block'], {}).get('options', []):
                    k1, k2 = norm(o['key']), norm(o['name'])
                    if subj and (subj in k2 or k1 in subj or subj in k1):
                        out += [r for _, r in e['rooms'].get(o['name'], [])]
                        break
            else:
                out += [r if r != 'いつもの教室' else cid[:2] + ' HB' for _, r in e['rooms']]
        return sorted(set(out))

    def base_room(name, i):
        lab = base[name][i]
        r, _ = room_of(lab, i)
        return r or from_class(lab, lambda c: student['classes'][c]['slots'][i])

    def lhr_fix(label, rooms):
        m = re.fullmatch(r'LHR([123][A-F])', norm(label))
        return [LHR_ROOM[m.group(1)]] if m and m.group(1) in LHR_ROOM else rooms

    base_rooms = {n: [lhr_fix(base[n][i], base_room(n, i)) for i in range(35)] for n in base}

    # ---- 特別時間割（教員版）
    special = {}
    names = {t['name'] for t in teachers}
    # 確定版の教室配当: {日付: {時限: {norm表記: [教室]}}}
    fixed = {}
    for f in sorted((ROOT / 'source' / 'kakutei').glob('*教室配当.pdf')):
        for d, per in parse_rooms(f).items():
            for p, lst in per.items():
                for room, lab in lst:
                    fixed.setdefault(d, {}).setdefault(p, {}).setdefault(norm(lab), []).append(room)
    for pdf in sorted((ROOT / 'source' / 'tokubetsu').glob('*.pdf')):
        for date, day in parse_pdf(pdf)[0].items():
            per = {}
            for n, cells in day['teachers'].items():
                n = RENAME.get(n, n)
                if n not in names:
                    print('  要確認: 特別時間割の名前が一覧にない', date, n)
                    continue
                labs = [cells.get(p, '') for p in range(1, day['periods'] + 1)]
                sc = student['special'].get(date, {}).get('classes', {})
                wd = datetime.date.fromisoformat(date).weekday()
                # 同じ曜日・時限のふだんの授業と同じ表記なら、その教室（確か）。ちがえばふだんの教室の目安
                fx = fixed.get(date, {})
                rs = [sorted(set(fx.get(p + 1, {}).get(norm(l), []))) or (room_of(l, wd * 7 + p)[0] if wd < 5 and norm(l) in by_slot.get(wd * 7 + p, {}) else [])
                      or room_of(l)[0] or from_class(l, lambda c, p=p: (sc.get(c) or [None] * 7)[p]) for p, l in enumerate(labs)]
                rs = [lhr_fix(l, r) for l, r in zip(labs, rs)]
                per[n] = {'l': labs, 'r': rs}
            special[date] = {'periods': day['periods'], 'note': day['note'], 'fixed': date in fixed, 't': per,
                             'src': '修学旅行特別時間割 ' + pdf.stem + ('（確定）' if date in fixed else '（第1案）')}

    # ---- 月間行事予定（全部の欄）
    staff_events = {}
    for pdf in sorted((ROOT / 'source').glob('gyouji_*.pdf')):
        m = re.search(r'gyouji_(\d+)', pdf.stem)
        month = int(m.group(1))
        year = 2026 if month >= 4 else 2027
        staff_events.update(gyouji_pdf.parse(pdf, year, month))

    payload = {
        'year': student['year'], 'updated': datetime.date.today().isoformat(),
        'teachers': teachers, 'base': base, 'baseRooms': base_rooms,
        'used': {str(i): sorted(v) for i, v in used.items()},
        'rooms': rooms_list, 'special': special, 'staffEvents': staff_events,
        'events': student['events'], 'jitei': student['jitei'], 'ready': student.get('ready', []),
    }

    pw = (ROOT / 'source' / 'staff_password.txt').read_text(encoding='utf-8').strip()
    salt, iv = os.urandom(16), os.urandom(12)
    key = PBKDF2HMAC(algorithm=hashes.SHA256(), length=32, salt=salt, iterations=ITER).derive(pw.encode())
    ct = AESGCM(key).encrypt(iv, json.dumps(payload, ensure_ascii=False, separators=(',', ':')).encode(), None)
    b64 = lambda b: base64.b64encode(b).decode()
    OUT.mkdir(exist_ok=True)
    (OUT / 'data.enc.js').write_text('window.ENC = ' + json.dumps({'salt': b64(salt), 'iv': b64(iv), 'iter': ITER, 'ct': b64(ct)}) + ';\n', encoding='utf-8')
    # 更新がすぐ届くよう、読みこむ側の index.html に版の印をつける（キャッシュ対策）
    import hashlib
    v = hashlib.sha1(ct).hexdigest()[:10]
    ih = OUT / 'index.html'
    ih.write_text(re.sub(r'<script src="data\.enc\.js[^"]*">', f'<script src="data.enc.js?v={v}">', ih.read_text(encoding='utf-8')), encoding='utf-8')
    print(f'data.enc.js を書き出しました（先生 {len(teachers)}・特別時間割 {len(special)} 日・月間行事 {len(staff_events)} 日・{len(ct)//1024} KB）')


if __name__ == '__main__':
    main()
