#!/usr/bin/env python3
"""時間割ナビのデータを作る。

  source/jikanwari.xlsx … 全体時間割（ルーム版教室入り・教室配当のシートを使う）
  予定.csv              … 特別時間割・行事・委員会・放課後の予定（手で書き足す）
  時程.csv              … 時程（通常・短縮など）

→ data.js を書き出す。index.html がこれを読む。
source/ は公開しない（先生の名前が入っているため）。
"""
import csv, json, re, sys, datetime, pathlib
import openpyxl
from collections import Counter
from invert import invert, ALIAS
from tokubetsu import parse_pdf
from kakutei import parse_rooms

ROOT = pathlib.Path(__file__).parent
DAYS = '月火水木金'

# ---------------------------------------------------------------- 文字の整え
FW = str.maketrans('０１２３４５６７８９ＡＢＣＤＥＦ', '0123456789ABCDEF')

def clean(v):
    if v is None:
        return ''
    s = str(v).replace('　', ' ').strip()
    s = re.sub(r'[ \t]+', ' ', s)
    return s

def oneline(s):
    return re.sub(r'\s*\n\s*', '', s).strip()

def norm(s):
    return oneline(s).translate(FW).replace(' ', '')

GROUP = {'①': 1, '②': 2, '③': 3}

# ---------------------------------------------------------------- クラス
# (id, 表示名, 学年, 科目行, 教室行, 系統) 系統は選択群の読みかえに使う
CLASSES = [
    ('1A', '1年A組', 1, 5, 6, ''), ('1B', '1年B組', 1, 7, 8, ''),
    ('1C', '1年C組', 1, 9, 10, ''), ('1D', '1年D組', 1, 11, 12, ''),
    ('1E', '1年E組', 1, 13, 14, ''),
    ('2A', '2年A組', 2, 15, 16, ''), ('2B', '2年B組', 2, 17, 18, ''),
    ('2C', '2年C組', 2, 19, 20, '2文'), ('2D', '2年D組', 2, 21, 22, '2文'),
    ('2E', '2年E組', 2, 23, 24, '2理'), ('2F', '2年F組', 2, 25, 26, '2理'),
    ('3A服飾', '3年A組（服飾）', 3, 27, 29, ''), ('3A食物', '3年A組（食物）', 3, 28, 29, ''),
    ('3B', '3年B組', 3, 30, 31, ''),
    ('3C', '3年C組', 3, 32, 33, '3文'), ('3D', '3年D組', 3, 34, 35, '3文'),
    ('3E', '3年E組', 3, 36, 37, '3理'), ('3F', '3年F組', 3, 38, 39, '3理'),
]
# 3D・3F の選択の欄は「科目：教室」の一覧で書かれているので、同じ系統のクラスから群を借りる
PARTNER = {'3D': '3C', '3F': '3E'}

# ---------------------------------------------------------------- 選択群
# key: 教室配当の表記とつき合わせる語。room: 教室配当に見つからないときの教室
def opt(name, room, key=None):
    return {'name': name, 'room': room, 'key': key or name}

ART = [opt('音楽', '音楽室'), opt('美術', '美術室'), opt('書道', '書道室')]
BLOCKS = {
    '芸術': {'label': '芸術', 'options': [opt('音楽Ⅰ', '音楽室', '音楽'), opt('美術Ⅰ', '美術室', '美術'), opt('書道Ⅰ', '書道室', '書道')]},
    '2文Ⅰ': {'label': '文系Ⅰ（Ⅰ選択）', 'tag': 'Ⅰ選択', 'options': [opt('日本史探究', '国語１'), opt('世界史探究', '国語２')]},
    '2理': {'label': '理系選択', 'tag': '理系選択', 'options': [opt('物理', '物理・地学室'), opt('生物', '生物室')]},
    '2福': {'label': '2年福祉選択', 'tag': '福選択', 'options': [opt('基礎国語', '2B HB'), opt('化学基礎', '物理・地学室')]},
    '3生': {'label': '3年生活選択', 'tag': '生選択', 'options': [opt('実用数学', '3A HB'), opt('生活芸術', '音楽室・書道室')]},
    '3福': {'label': '3年福祉選択', 'tag': '福選択', 'options': [opt('実用数学', 'ゼミ１'), opt('社会福祉探究', '3B HBほか')]},
    'ア': {'options': [opt('世界史特講', '地歴３'), opt('日本史特講', '地歴２'), opt('地理特講', 'HB4-1'),
                      opt('政治・経済特講', 'ゼミ１', '政経特講'), opt('文系数学特講ⅡＢＣ', 'ゼミ２', '特講ⅡＢＣ')]},
    'イ': {'options': [opt('世界史実践', '地歴３'), opt('日本史実践', '地歴２'),
                      opt('文系数学特講ⅠＡ', 'ゼミ２', '特講ⅠＡ'), opt('進路実践', '自習室')]},
    'ウ': {'options': [opt('地理実践', 'HB4-1'), opt('公民実践', '多目的２・多目的３'),
                      opt('総合音楽', '音楽室'), opt('総合書道', '書道室'), opt('総合美術', '美術室'), opt('進路実践', '自習室')]},
    'エ': {'options': [opt('文系数学特講α', 'ゼミ２', '特講α'), opt('英語実践', '国語３・多目的２・多目的３')]},
    'オ': {'options': [opt('数学実践', '数学２'), opt('英語特講', 'ゼミ１'), opt('実用英語', 'ゼミ２'), opt('進路実践', '自習室')]},
    'カ': {'options': [opt('情報実践', '情報処理室'), opt('国語表現研究', '地歴２', '表現研究'),
                      opt('実用英語', 'ゼミ２'), opt('進路実践', '自習室')]},
    'キ': {'options': [opt('理科基礎実践', '理科教室', '理科基礎'), opt('生涯スポーツ', '体育施設'),
                      opt('総合音楽', '音楽室'), opt('総合書道', '書道室'), opt('総合美術', '美術室')]},
    'サ': {'options': [opt('数学Ⅲ', '数学１・数学２'), opt('理系数学特講ⅡＢＣ', 'ゼミ２', '特講ⅡＢＣ')]},
    'シ': {'options': [opt('理系数学特講ＢＣ', '数学１・数学２', '特講ＢＣ'), opt('理系数学特講ⅠＡ', 'ゼミ２', '特講ⅠＡ')]},
    'ス': {'options': [opt('化学特講', '理科教室・化学室'), opt('物理実践', '物理・地学室'), opt('生物実践', '生物室')]},
    'セ': {'options': [opt('物理特講', '物理・地学室'), opt('生物特講', '生物室'), opt('化学実践', '理科教室')]},
    'ソ': {'options': [opt('国語実践', '地歴２・HB5-2'), opt('総合音楽', '音楽室'), opt('総合書道', '書道室'),
                      opt('総合美術', '美術室'), opt('進路実践', '自習室')]},
    'タ': {'options': [opt('理系数学特講Ⅲ', '数学１・数学２', '数学特講Ⅲ'), opt('英語実践', 'ゼミ１')]},
    'チ': {'options': [opt('数学実践', '数学３'), opt('英語特講', '多目的１')]},
    'ツ': {'options': [opt('倫理', 'ゼミ１'), opt('政治・経済', '地歴３', '政治経済'),
                      opt('社会実践（三総合）', 'HB4-2', '三総合'), opt('社会実践（地理）', 'HB4-1', '(地理'),
                      opt('国語表現研究', '地歴２', '表現研究'), opt('実用英語', 'ゼミ２'), opt('進路実践', '自習室')]},
    'テ': {'options': [opt('情報実践', '情報処理室'), opt('生涯スポーツ', '体育施設'),
                      opt('総合音楽', '音楽室'), opt('総合書道', '書道室'), opt('総合美術', '美術室')]},
}
for k, b in BLOCKS.items():
    b.setdefault('label', k + '選択')
    b.setdefault('tag', k + '選択')


def block_of(subj, cls_course, cls_id):
    """ルーム版の科目欄から選択群を返す。選択でなければ None。"""
    s = norm(subj)
    if s == '芸術':
        return '芸術'
    if 'Ⅰ選択' in s:
        return '2文Ⅰ'
    if s == '理系選択':
        return '2理'
    if '福祉選択' in s:
        return '2福' if cls_id == '2B' else '3福'
    if '生活選択' in s:
        return '3生'
    m = re.fullmatch(r'([ア-ン]{1,2})選択', s)
    if m:
        letters = m.group(1)
        if len(letters) == 1:
            return letters
        return letters[0] if cls_course == '3文' else letters[1]  # キテ・ウソ・カツ
    return None


# ---------------------------------------------------------------- Excel を読む
def filled(ws):
    g = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                g[(c.row, c.column)] = c.value
    for m in ws.merged_cells.ranges:
        v = g.get((m.min_row, m.min_col))
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                g[(r, c)] = v
    return g


def load_room_sheet(wb):
    """教室配当: {slot: [(教室, 表記)]}"""
    ws = wb['教室配当']
    g = filled(ws)
    by_slot = {}
    for r in range(6, ws.max_row + 1):
        room = oneline(clean(g.get((r, 2))))
        if not room:
            continue
        for i in range(35):
            lab = clean(g.get((r, 3 + i)))
            if lab:
                by_slot.setdefault(i, []).append((room, lab))
    return by_slot


def split_rooms(text):
    return [oneline(x) for x in re.split(r'\n', text) if oneline(x)]


def parse_pairs(text):
    """「日本史探究：国語１\n世界史探究：国語２」→ [(科目, 教室)]"""
    out = []
    for line in text.split('\n'):
        line = line.strip()
        if '：' in line:
            a, b = line.split('：', 1)
            out.append((a.strip(), b.strip()))
        elif out and line:
            out[-1] = (out[-1][0], out[-1][1] + '・' + line)
    return out


def tidy_room(r):
    r = r.replace('Ｈ', 'H').replace('Ｂ', 'B').replace('体育', '体育施設').replace('体育施設施設', '体育施設')
    r = r.replace('物理地学室', '物理・地学室').replace('国語1', '国語１').replace('，', '・').replace(',', '・')
    m = re.fullmatch(r'(\d[A-F])HB(.*)', r.translate(FW))
    if m:
        r = f'{m.group(1)} HB{m.group(2)}'
    return r


def lookup_assigned(by_slot, slot, test):
    """教室配当でその時間に test(表記) を満たす教室を [(班, 教室)] で返す"""
    hits = []
    for room, lab in by_slot.get(slot, []):
        n = norm(lab)
        if '－' in lab or '-' in n and re.search(r'\d-\d', n):
            continue  # 附属中（３－１ など）
        if test(n):
            grp = next((GROUP[ch] for ch in lab if ch in GROUP), 0)
            hits.append((grp, tidy_room(room)))
    hits.sort()
    return hits


def build_timetable(wb, warn):
    ws = wb['ルーム版教室入り']
    g = filled(ws)
    by_slot = load_room_sheet(wb)
    classes = {}
    used_blocks = set()
    for cid, label, grade, srow, rrow, course in CLASSES:
        code = re.sub(r'[^0-9A-F]', '', cid.translate(FW))[:2]  # '3A服飾' → '3A'
        slots = []
        for i in range(35):
            subj_raw = clean(g.get((srow, 2 + i)))
            room_raw = clean(g.get((rrow, 2 + i)))
            subj = oneline(subj_raw)
            if not subj:
                slots.append(None)
                continue
            # 3D・3F の「科目：教室」一覧 → 相方クラスの群
            if '：' in subj_raw and cid in PARTNER:
                p = PARTNER[cid]
                pc = next(c for c in CLASSES if c[0] == p)
                subj_raw = clean(g.get((pc[3], 2 + i)))
                subj = oneline(subj_raw)
            blk = block_of(subj, course, cid)
            if blk:
                used_blocks.add(blk)
                b = BLOCKS[blk]
                tag = norm(b['tag'])
                rooms = {}
                for o in b['options']:
                    key = norm(o['key'])
                    hits = lookup_assigned(by_slot, i, lambda n: key in n and (tag in n or blk == '芸術'))
                    rooms[o['name']] = [list(h) for h in hits] if hits else [[0, o['room']]]
                slots.append({'block': blk, 'rooms': rooms})
                continue
            # Ⅱ選択（倫理・政経）は全員同じ
            m = re.match(r'Ⅱ選択(.+)', subj)
            note = ''
            if m:
                subj, note = m.group(1), 'Ⅱ選択'
            split = '(分割)' in subj
            subj = subj.replace('(分割)', '').strip()
            subj = re.sub(r'(現代の|科学と)(国語|人間生活)', r'\1\2', subj)
            subj = ALIAS.get(subj, subj)
            # 教室: まず教室配当、なければルーム版の教室欄
            key = norm(subj)[:2]
            hits = lookup_assigned(by_slot, i, lambda n: code in n and '選択' not in n and n.startswith(key))
            if not hits:
                hits = lookup_assigned(by_slot, i, lambda n: code in n and '選択' not in n and 'LHR' not in n and '総合' not in n)
                if len(hits) > 1 and len({h[0] for h in hits}) == 1:
                    hits = []
            rooms = [list(h) for h in hits]
            if '：' in room_raw:
                rooms = []  # 「科目：教室」の書き分けがあればそちらを優先
            if not rooms and room_raw:
                if '：' in room_raw:
                    pairs = parse_pairs(room_raw)
                    pick = [p for p in pairs if norm(p[0])[:2] == key[:2]]
                    rooms = [[0, tidy_room(pick[0][1])]] if pick else [[0, tidy_room(p[1])] for p in pairs]
                else:
                    rs = split_rooms(room_raw)
                    rooms = [[k + 1 if len(rs) > 1 else 0, tidy_room(r)] for k, r in enumerate(rs)]
            if rooms and room_raw and '：' not in room_raw:
                a = sorted(tidy_room(r) for r in split_rooms(room_raw))
                b2 = sorted(r for _, r in rooms)
                if a != b2:
                    warn.append(f'{cid} {DAYS[i//7]}{i%7+1} {subj}: 教室配当={b2} ルーム版={a}')
            if not rooms:
                if subj in ('体育', '生涯スポーツ'):
                    rooms = [[0, '体育施設']]
                elif subj == '総合':
                    rooms = [[0, '学年の指示どおり']]
                else:
                    rooms = [[0, 'いつもの教室']]
            if len(rooms) > 1 and all(r[0] for r in rooms):
                split = True
            ent = {'s': subj, 'rooms': rooms}
            if note:
                ent['note'] = note
            if split and len(rooms) > 1:
                ent['split'] = True
            slots.append(ent)
        # 3A 食物は服飾と違う欄だけ書かれている
        classes[cid] = {'label': label, 'grade': grade, 'course': course, 'slots': slots}
    blocks = {k: {'label': v['label'], 'options': [o['name'] for o in v['options']]} for k, v in BLOCKS.items() if k in used_blocks}
    return classes, blocks


# ---------------------------------------------------------------- 特別時間割
def bare(s):
    return re.sub(r'[ⅠⅡⅢ]|\|.*', '', s)


def build_special(classes, warn):
    """source/tokubetsu/*.pdf（教員版）→ {日付: {note, periods, classes: {id: [7コマ]}}}
    教室は載っていないので、ふだんの時間割で同じ科目が使う教室（いちばん多いもの）をあてる。"""
    def room_for(cid, subj):
        c = classes[cid]
        for test in (lambda s: s == subj, lambda s: bare(s) == bare(subj), lambda s: s[:2] == subj[:2]):
            cnt = Counter(json.dumps(e['rooms'], ensure_ascii=False) for e in c['slots'] if e and 'block' not in e and test(e['s']))
            if cnt:
                (best, n), total = cnt.most_common(1)[0], sum(cnt.values())
                return json.loads(best), n < total
        if subj in ('体育', '生涯スポーツ'):
            return [[0, '体育施設']], False
        if subj == '総合':
            return [[0, '学年の指示どおり']], False
        return [[0, 'いつもの教室']], False

    def block_rooms(cid, blk):
        cnt = Counter(json.dumps(e['rooms'], ensure_ascii=False) for e in classes[cid]['slots'] if e and e.get('block') == blk)
        return json.loads(cnt.most_common(1)[0][0]) if cnt else {}

    # 確定版の教室配当（source/kakutei/*_教室配当.pdf）があれば、その日はその教室を使う
    fixed = {}
    for f in sorted((ROOT / 'source' / 'kakutei').glob('*教室配当.pdf')):
        fixed.update(parse_rooms(f))

    def fixed_rooms(date, p, test):
        return [list(h) for h in lookup_assigned({0: fixed[date].get(p, [])}, 0, test)]

    special = {}
    for pdf in sorted((ROOT / 'source' / 'tokubetsu').glob('*.pdf')):
        for date, day in parse_pdf(pdf)[0].items():
            inv = invert(day['labels'])
            per = {}
            for cid in classes:
                slots = [None] * 7
                for p, ent in inv.get(cid, {}).items():
                    if p > day['periods']:
                        continue
                    if ent[0] == 'block':
                        if ent[1] not in BLOCKS:
                            continue
                        wd = datetime.date.fromisoformat(date).weekday()
                        same = classes[cid]['slots'][wd * 7 + p - 1] if wd < 5 else None
                        rooms = same['rooms'] if same and same.get('block') == ent[1] else block_rooms(cid, ent[1])
                        if date in fixed:
                            tag = norm(BLOCKS[ent[1]]['tag'])
                            fr = {}
                            for o in BLOCKS[ent[1]]['options']:
                                key = norm(o['key'])
                                h = fixed_rooms(date, p, lambda n: key in n and (tag in n or ent[1] == '芸術'))
                                fr[o['name']] = h or rooms.get(o['name'], [[0, o['room']]])
                            rooms = fr
                        slots[p - 1] = {'block': ent[1], 'rooms': rooms}
                    else:
                        names = [x.split('|')[0] for x in ent[1]]
                        # 表記ゆれ（英コミュⅡ/Ⅲ など）はふだんの時間割の名前にそろえる
                        base_names = {e['s'] for e in classes[cid]['slots'] if e and 'block' not in e}
                        names = list(dict.fromkeys(next((b for b in base_names if bare(b) == bare(n)), n) if n not in base_names else n for n in names))
                        note = 'Ⅱ選択' if any('|Ⅱ選択' in x for x in ent[1]) else ''
                        # 同じ曜日・時限のふだんの授業と同じなら、その教室（確か）
                        wd = datetime.date.fromisoformat(date).weekday()
                        same = classes[cid]['slots'][wd * 7 + p - 1] if wd < 5 else None
                        code = re.sub(r'[^0-9A-F]', '', cid.translate(FW))[:2]
                        key = norm(names[0])[:2]
                        hit = fixed_rooms(date, p, lambda n: code in n and '選択' not in n and n.startswith(key)) if date in fixed else []
                        if not hit and date in fixed and names[0] in ('ＬＨＲ', 'LHR'):
                            hit = fixed_rooms(date, p, lambda n: code in n and 'LHR' in n)
                        if hit:
                            rooms, varies = hit, False
                        elif same and 'block' not in same and same['s'] == names[0]:
                            rooms, varies = same['rooms'], False
                        else:
                            rooms, varies = room_for(cid, names[0])
                        e = {'s': ' / '.join(names), 'rooms': rooms}
                        if note:
                            e['note'] = note
                        if len(names) > 1 and e['s'] not in base_names:
                            e['conflict'] = True
                            warn.append(f'特別 {date} {cid} {p}限: 案で重なり {e["s"]}')
                        slots[p - 1] = e
                # 途中のコマが空いているのは読みとり漏れかもしれない
                filled_p = [i for i, x in enumerate(slots) if x]
                if filled_p:
                    gaps = [i + 1 for i in range(filled_p[0], filled_p[-1]) if not slots[i]]
                    if gaps:
                        warn.append(f'特別 {date} {cid}: {gaps}限が空き')
                per[cid] = slots
            special[date] = {'periods': day['periods'], 'classes': per, 'fixed': date in fixed,
                             'src': '修学旅行特別時間割 ' + pdf.stem + ('（確定）' if date in fixed else '（第1案）')}
    return special


# ---------------------------------------------------------------- 予定・時程
def read_csv(name):
    p = ROOT / name
    if not p.exists():
        return []
    with open(p, encoding='utf-8-sig') as f:
        rows = [r for r in csv.DictReader(f)]
    return [{k.strip(): (v or '').strip() for k, v in r.items() if k} for r in rows if any((v or '').strip() for v in r.values())]


# 資料どうしで食い違う教室を、本人に確かめた答えで直す（教室配当の LHR は 1C・1D が入れ違っている。本人確認 2026-10-10）
LHR_ROOM = {'1C': '国語３', '1D': '多目的２'}


def fix_lhr(slots, cid):
    for e in slots:
        if e and 'block' not in e and e['s'] in ('ＬＨＲ', 'LHR') and cid in LHR_ROOM:
            e['rooms'] = [[0, LHR_ROOM[cid]]]


def main():
    wb = openpyxl.load_workbook(ROOT / 'source' / 'jikanwari.xlsx', data_only=True)
    warn = []
    classes, blocks = build_timetable(wb, warn)
    warn[:] = [w for w in warn if not any(w.startswith(f'{c} 金7 ＬＨＲ') for c in LHR_ROOM)]
    special = build_special(classes, warn)
    for cid, c in classes.items():
        fix_lhr(c['slots'], cid)
    for day in special.values():
        for cid, slots in day['classes'].items():
            fix_lhr(slots, cid)
    jitei = {}
    for r in read_csv('時程.csv'):
        jitei.setdefault(r['時程名'], []).append({'p': r['時限'], 'start': r['開始'], 'end': r['終了']})
    events = []
    for r in read_csv('予定.csv'):
        if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', r.get('日付', '')):
            warn.append(f'予定.csv の日付が読めない: {r}')
            continue
        events.append({k: v for k, v in r.items() if v})
    ready = []
    p_ready = ROOT / '表示する月.txt'
    if p_ready.exists():
        for line in p_ready.read_text(encoding='utf-8').splitlines():
            line = line.split('#')[0].strip()
            if re.fullmatch(r'\d{4}-\d{2}', line):
                ready.append(line)
            elif line:
                warn.append(f'表示する月.txt の行が読めない: {line}')
    data = {
        'year': '2026年度',
        'updated': datetime.date.today().isoformat(),
        'classes': classes,
        'blocks': blocks,
        'jitei': jitei,
        'events': events,
        'special': special,
        'ready': sorted(ready),
    }
    js = 'window.DATA = ' + json.dumps(data, ensure_ascii=False, separators=(',', ':')) + ';\n'
    (ROOT / 'data.js').write_text(js, encoding='utf-8')
    # 更新がすぐ届くよう、index.html の読みこみに版の印をつける（キャッシュ対策）
    import hashlib
    v = hashlib.sha1(js.encode()).hexdigest()[:10]
    ih = ROOT / 'index.html'
    ih.write_text(re.sub(r'<script src="data\.js[^"]*">', f'<script src="data.js?v={v}">', ih.read_text(encoding='utf-8')), encoding='utf-8')
    print(f'data.js を書き出しました（クラス {len(classes)}・選択群 {len(blocks)}・予定 {len(events)} 件・特別時間割 {len(special)} 日・表示する月 {", ".join(ready)}）')
    for w in warn:
        print('  要確認:', w)


def show(ids):
    """確認用: python3 build.py --show 3C 2E"""
    d = json.loads((ROOT / 'data.js').read_text(encoding='utf-8')[len('window.DATA = '):-2])
    for cid in ids or d['classes']:
        c = d['classes'][cid]
        print('=====', cid, c['label'])
        for di in range(5):
            row = []
            for p in range(7):
                e = c['slots'][di * 7 + p]
                if not e:
                    row.append('－')
                elif 'block' in e:
                    row.append(f"[{e['block']}]")
                else:
                    row.append(e['s'] + '@' + '/'.join((f'{g}:' if g else '') + x for g, x in e['rooms']))
            print(DAYS[di], ' | '.join(row))


if __name__ == '__main__':
    if '--show' in sys.argv:
        show([a for a in sys.argv[1:] if a != '--show'])
    else:
        main()
